from pypdf import PdfReader
import docx
from bs4 import BeautifulSoup
import openpyxl

def load_pdf(path):
    reader = PdfReader(path)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def load_docx(path):
    doc = docx.Document(path)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text

def load_txt(path):
    with open(path, "r", encoding="utf-8") as file:
        return file.read()

def load_html(path):
    with open(path, "r", encoding="utf-8") as file:
        soup = BeautifulSoup(file, "html.parser")
        return soup.get_text(separator="\n")

def load_excel(path):
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0][:12]  # ✅ Limit header to first 12 columns
    data_rows = [row[:12] for row in rows[1:]]  # ✅ Limit each row to first 12 columns

    return header, data_rows

def chunk_excel_rows(header, data_rows):
    chunks = []
    metadatas = []

    for idx, row in enumerate(data_rows):
        metadata = {
            f"col_{header[i]}": str(row[i]) if row[i] is not None else ""
            for i in range(min(9, len(row)))
        }
        content_cells = row[9:] if len(row) > 9 else []
        content = " | ".join(str(cell) for cell in content_cells if cell is not None)

        # Split at pipe to isolate matchable content
        question_part = content.split("|")[0].strip() if "|" in content else content
        full_content = content if content else "No content"

        # Use question part for embedding, full content for context
        chunks.append({"embed": question_part, "full": full_content})
        metadatas.append(metadata)

    return chunks, metadatas

def load_document(path):
    if path.endswith(".pdf"):
        return load_pdf(path)
    elif path.endswith(".docx"):
        return load_docx(path)
    elif path.endswith(".html") or path.endswith(".htm"):
        return load_html(path)
    elif path.endswith(".txt"):
        return load_txt(path)
    elif path.endswith(".xlsx"):
        header, data_rows = load_excel(path)
        return chunk_excel_rows(header, data_rows)
    else:
        raise ValueError("Unsupported file format")
